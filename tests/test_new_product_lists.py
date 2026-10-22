"""New-product list: scan, name + supplier + category + quantity, many on one list, save -> Excel."""
import pytest
from sqlalchemy import select

from app.models import NewProductList, Supplier
from app.services import new_product_lists as npl
from app.services import new_products
from tests.test_web import add_user, client, csrf_of, login, seeded  # noqa: F401 (fixtures)


@pytest.fixture(autouse=True)
def _no_storehub_or_internet(monkeypatch):
    monkeypatch.setattr(new_products, "refresh_products", lambda: False)
    monkeypatch.setattr(new_products, "web_lookup", lambda db, b, timeout=4.0: None)


def _setup(sync, fake, session, user="np_user"):
    seeded(sync, fake, session)
    session.add(Supplier(name="Acme Foods"))
    session.commit()
    add_user(session, user, "staff", locs=[2])


def test_options_come_from_existing_suppliers_and_categories(sync, fake, session):
    _setup(sync, fake, session)
    assert "Acme Foods" in npl.suppliers(session) and {"Drinks", "Snacks"} <= set(npl.categories(session))


def test_add_many_then_save_makes_the_rpa_excel(cfg, sync, fake, session, client):
    _setup(sync, fake, session)
    login(client, "np_user")
    page = client.get("/new-product")
    assert page.status_code == 200 and "Scan the barcode" in page.text and "Acme Foods" in page.text and '"Drinks"' in page.text and 'id="np-supplier"' in page.text
    csrf = csrf_of(client, "/new-product")
    ok = client.post("/new-product/add", data={"csrf": csrf, "barcode": "880011", "name": "Staff Snack 50g", "supplier": "Acme Foods", "category": "Snacks", "qty": "12"})
    assert ok.status_code == 200 and ok.json()["total_lines"] == 1
    r = client.post("/new-product/add", data={"csrf": csrf, "barcode": "880022", "name": "Staff Drink", "supplier": "Acme Foods", "category": "Drinks", "qty": "6"}).json()
    assert r["ok"] and r["total_lines"] == 2 and r["total_qty"] == 18
    save = client.post("/new-product/save", data={"csrf": csrf})
    assert save.status_code == 303 and "err=" not in save.headers["location"]
    assert "/new-product/lists/" in save.headers["location"]  # after Save the list itself opens, with its StoreHub status
    import re
    hid = lambda url: re.search(r"<div id=.np-saved-wrap.([^>]*)>", client.get(url).text).group(1)
    assert "hidden" in hid("/new-product") and "hidden" not in hid("/new-product?saved=1")
    session.expire_all()
    lst = session.scalar(select(NewProductList))
    assert lst.status == "saved" and lst.excel_path
    xl = client.get(f"/new-product/lists/{lst.id}/excel")
    assert xl.status_code == 200 and xl.content[:2] == b"PK"
    from openpyxl import load_workbook
    import io
    rows = list(load_workbook(io.BytesIO(xl.content)).active.iter_rows(values_only=True))
    assert list(rows[0]) == new_products.EXCEL_HEADERS and len(rows) == 3
    assert rows[1][0] == "Staff Snack 50g" and rows[1][2] == "880011" and rows[1][3] == "Snacks" and rows[1][4] == "Acme Foods" and rows[1][7] == 12 and rows[2][2] == "880022"
    assert (cfg.data_dir / lst.excel_path).parent == (cfg.data_dir / new_products.RPA_DIR)  # the folder the RPA watches
    assert "Staff Snack 50g" in client.get(f"/new-product/lists/{lst.id}").text
    # the saved list can no longer be changed, and the next scan starts a new list
    assert client.post("/new-product/add", data={"csrf": csrf, "barcode": "880033", "name": "N", "supplier": "Acme Foods", "category": "Drinks", "qty": "1"}).json()["total_lines"] == 1
    assert client.post(f"/new-product/lists/{lst.id}/delete", data={"csrf": csrf}).status_code == 303
    assert client.get(f"/new-product/lists/{lst.id}/excel").status_code == 404


def test_checks_refuse_bad_lines(sync, fake, session, client):
    _setup(sync, fake, session)
    login(client, "np_user")
    csrf = csrf_of(client, "/new-product")
    good = {"csrf": csrf, "barcode": "880011", "name": "Staff Snack", "supplier": "Acme Foods", "category": "Snacks", "qty": "3"}
    for bad, why in (({"barcode": "111"}, "Already in StoreHub"), ({"name": " "}, "product name"), ({"supplier": "Nobody Ltd"}, "supplier"),
                     ({"category": "Made Up"}, "category"), ({"qty": "0"}, "above zero"), ({"qty": "1.5"}, "whole number"), ({"qty": ""}, "quantity")):
        r = client.post("/new-product/add", data={**good, **bad})
        assert r.status_code == 400 and why in r.json()["error"], (bad, r.text)
    assert client.post("/new-product/add", data=good).json()["ok"]
    assert "already on the list" in client.post("/new-product/add", data=good).json()["error"]
    assert client.post("/new-product/save", data={"csrf": csrf}).status_code == 303
    assert "err=" in client.post("/new-product/save", data={"csrf": csrf}).headers["location"]  # nothing left to save: a fresh empty list


def test_remove_a_line_and_check_endpoint(sync, fake, session, client):
    _setup(sync, fake, session)
    login(client, "np_user")
    csrf = csrf_of(client, "/new-product")
    rows = client.post("/new-product/add", data={"csrf": csrf, "barcode": "880011", "name": "A", "supplier": "Acme Foods", "category": "Snacks", "qty": "1"}).json()
    assert client.get("/new-product/check?barcode=111").json() == {"known": "Cola 500ml"}
    assert client.get("/new-product/check?barcode=999999").json() == {"known": None}
    left = client.post(f"/new-product/line/{rows['lines'][0]['id']}/remove", data={"csrf": csrf}).json()
    assert left["ok"] and left["total_lines"] == 0


def test_lists_are_private_to_their_owner(sync, fake, session, client):
    _setup(sync, fake, session)
    add_user(session, "np_other", "staff", locs=[2])
    login(client, "np_user")
    csrf = csrf_of(client, "/new-product")
    client.post("/new-product/add", data={"csrf": csrf, "barcode": "880011", "name": "A", "supplier": "Acme Foods", "category": "Snacks", "qty": "1"})
    client.post("/new-product/save", data={"csrf": csrf})
    lid = session.scalar(select(NewProductList.id))
    client.cookies.clear()
    login(client, "np_other")
    assert client.get(f"/new-product/lists/{lid}/excel").status_code == 404


def test_discard_throws_away_the_whole_open_list(sync, fake, session, client):
    _setup(sync, fake, session)
    login(client, "np_user")
    csrf = csrf_of(client, "/new-product")
    for bc in ("880011", "880022"):
        assert client.post("/new-product/add", data={"csrf": csrf, "barcode": bc, "name": "A" + bc, "supplier": "Acme Foods", "category": "Snacks", "qty": "2"}).json()["ok"]
    r = client.post("/new-product/discard", data={"csrf": csrf})
    assert r.status_code == 303 and "err=" not in r.headers["location"]
    session.expire_all()
    assert session.scalars(select(NewProductList)).all() == []
    assert client.post("/new-product/discard", data={"csrf": csrf}).status_code == 303  # nothing to discard is fine
    assert client.post("/new-product/add", data={"csrf": csrf, "barcode": "880011", "name": "Again", "supplier": "Acme Foods", "category": "Snacks", "qty": "1"}).json()["total_lines"] == 1
