"""The store sales report: an Excel file in the layout of StoreHub's own sales report (a sheet per month, Monthly Summary, All Sales)."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Location, SalesDay, SalesTx
from app.services import sales_report as sr
from app.security import hash_password
from app.timeutil import now_utc
from tests.test_ai import ready  # noqa: F401
from tests.test_web import PW, client, login  # noqa: F401


def test_title_and_file_name():
    class L:
        name = "The Daily Centro - Tabunoc"
    assert sr.title_of(L) == "TDC-TABUNOC" and sr.file_name(L) == "Sales_Report_TDC-TABUNOC.xlsx"
    L.name = "The Daily Centro - Bantayan"
    assert sr.title_of(L) == "TDC-BANTAYAN"
    assert sr.label(date(2026, 8, 26)) == "26 Aug 2026 (Wed)" and sr.label(date(2026, 10, 1)) == "1 Oct 2026 (Thu)"


def test_the_workbook_has_the_storehub_report_layout_and_numbers(session, sync, fake, client):
    boss = ready(sync, fake, session)
    boss.password_hash = hash_password(PW)
    session.commit()
    loc = session.scalar(select(Location).where(Location.type == "store"))
    # a known day, as in the Tabunoc report: total 2,199.00, tax 235.56, discount 1.79 -> net 1,963.44
    d = date(2026, 9, 29)
    for old in session.scalars(select(SalesDay).where(SalesDay.location_id == loc.id, SalesDay.business_date.in_([d, d + timedelta(days=1)]))):
        session.delete(old)
    session.flush()
    session.add(SalesDay(location_id=loc.id, business_date=d, tx_count=40, gross_sales=2199.0, returns_total=0.0, net_sales=2199.0, tax=235.56, cash_sales=1958.0, other_sales=241.0))
    session.add(SalesTx(refid="r1", invoice_no="1", location_id=loc.id, tx_type="Sale", tx_time=now_utc(), business_date=d, total=2199.0, subtotal=1965.23, tax=235.56, discount=1.79,
                        rounded=0.0, service_charge=0.0, cancelled=False, raw_hash="x"))
    session.commit()
    login(client, "boss")
    r = client.get(f"/sales/report.xlsx?loc={loc.id}")
    assert r.status_code == 200 and "Sales_Report_TDC-" in r.headers["content-disposition"] and r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[-2:] == ["Monthly Summary", "All Sales"] and "Sep 2026" in wb.sheetnames
    ws = wb["All Sales"]
    assert [c.value for c in ws[2]] == sr.HEADERS and ws["A1"].value == sr.title_of(loc) and ws.freeze_panes == "B3"
    row = next(r_ for r_ in ws.iter_rows(min_row=3, values_only=True) if r_[0] == "29 Sep 2026 (Tue)")
    # Total Sales, Transactions, Discount, Discount %, Manual Discount, Manual %, Tax, Rounding, Shipping, Service, GP, GP%, Returned, Net, Average
    assert row[1:11] == (2199.0, 40, 1.79, 0.09, 1.79, 0.08, 235.56, 0, 0, 0)
    assert row[13] == 0 and row[14] == 1963.44 and row[15] == 49.09  # Net Sales = total - tax; average = net / transactions
    assert row[11] is None and row[12] is None  # no cost history for that day: Gross Profit is left blank, not guessed
    month = wb["Sep 2026"]
    last = month.max_row
    assert month.cell(row=last, column=1).value == "TOTAL" and month.cell(row=last, column=2).value.startswith("=SUM(B3:B")
    summary = wb["Monthly Summary"]
    assert summary["A2"].value.startswith("Monthly Summary") and "'All Sales'!B" in summary["C5"].value


def test_a_store_without_sales_gets_a_message_not_an_empty_file(session, sync, fake, client):
    boss = ready(sync, fake, session)
    boss.password_hash = hash_password(PW)
    session.commit()
    empty = Location(storehub_id="EMPTY", name="The Daily Centro - Empty", type="store", monitored=True)
    session.add(empty)
    session.commit()
    login(client, "boss")
    r = client.get(f"/sales/report.xlsx?loc={empty.id}", follow_redirects=False)
    assert r.status_code == 303 and "err=" in r.headers["location"]


def test_gross_profit_comes_from_storehubs_report_and_fills_every_day(session, sync, fake, client, monkeypatch):
    from app import robot
    from app.services import sales_profit
    boss = ready(sync, fake, session)
    boss.password_hash = hash_password(PW)
    session.commit()
    loc = session.scalar(select(Location).where(Location.type == "store", Location.monitored.is_(True)))
    days = list(session.scalars(select(SalesDay).where(SalesDay.location_id == loc.id).order_by(SalesDay.business_date)))
    assert days and all(d.gross_profit is None for d in days)
    # StoreHub sends a day as its local midnight in UTC: 16:00Z the evening before
    assert sales_profit.local_date("2026-10-09T16:00:00.000Z") == date(2026, 10, 10)
    day = days[0]
    stamp = f"{(day.business_date - timedelta(days=1)).isoformat()}T16:00:00.000Z"
    monkeypatch.setattr(robot, "configured", lambda: True)
    asked = {}

    def fake_fetch(names, start, end):
        asked.update(names=names, start=start, end=end)
        return {loc.name: [{"date": stamp, "grossProfit": 1724.13}, {"date": "2020-01-01T16:00:00.000Z", "grossProfit": 5}]}
    out = sales_profit.run(session, fetch=fake_fetch)
    assert out["ran"] and out["ok"] and out["days"] == 1 and loc.name in asked["names"]
    session.expire_all()
    assert session.get(SalesDay, (loc.id, day.business_date)).gross_profit == 1724.13
    assert sales_profit.due(session) is False or out["stores"] < len(sales_profit.stores(session))  # not asked again straight away unless a store was missing
    login(client, "boss")
    wb = load_workbook(io.BytesIO(client.get(f"/sales/report.xlsx?loc={loc.id}").content))
    ws = wb["All Sales"]
    row = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] == sr.label(day.business_date))
    assert row[11] == 1724.13 and row[12] is not None  # StoreHub's Gross Profit, and its percentage of Net Sales
