"""New-product registration: validation, Lark mirror, RPA Excel export, and the web flow."""
import json

import httpx
import pytest
from sqlalchemy import func, select

from app import audit as audit_mod
from app.config import Config
from app.lark import LarkClient
from app.models import AuditEvent, NewProduct
from app.services import new_products

_REAL_WEB_LOOKUP = new_products.web_lookup
from tests.test_web import add_user, client, csrf_of, login, seeded  # noqa: F401 (fixtures)


@pytest.fixture(autouse=True)
def _no_storehub_or_internet(monkeypatch):
    """The page asks StoreHub for the product list and the internet for the barcode: neither is ever reached from a test."""
    monkeypatch.setattr(new_products, "refresh_products", lambda: False)
    monkeypatch.setattr(new_products, "web_lookup", lambda db, b, timeout=4.0: None)


class FakeLarkBase:
    """Just enough of the Lark Open API for the New Product table. fields: list of (name, type) — type 20 is a computed column."""
    def __init__(self, fields, fail=False):
        self.fields, self.fail, self.rows = list(fields), fail, []

    def handler(self, request):
        path = request.url.path
        if path.endswith("tenant_access_token/internal"):
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "t", "expire": 7200})
        if self.fail:
            return httpx.Response(200, json={"code": 999, "msg": "boom"})
        if path.endswith("/fields") and request.method == "GET":
            return httpx.Response(200, json={"code": 0, "data": {"items": [{"field_name": n, "type": t} for n, t in self.fields], "has_more": False}})
        if path.endswith("/fields") and request.method == "POST":
            body = json.loads(request.content)
            self.fields.append((body["field_name"], body["type"]))
            return httpx.Response(200, json={"code": 0, "data": {}})
        if path.endswith("batch_create"):
            body = json.loads(request.content)
            self.rows.append(body["records"][0]["fields"])
            return httpx.Response(200, json={"code": 0, "data": {"records": [{"record_id": f"rec{len(self.rows)}"}]}})
        return httpx.Response(404, json={"code": 1, "msg": "nope"})


def fake_client(cfg, fields, fail=False):
    f = FakeLarkBase(fields, fail)
    c = Config(data_dir=cfg.data_dir, lark_app_id="id", lark_app_secret="sec", lark_base_token=new_products.LARK_BASE_TOKEN, scheduler_enabled=False)
    return f, LarkClient(c, transport=httpx.MockTransport(f.handler), sleep=lambda s: None)


def test_register_writes_lark_row_excel_and_audit(cfg, session):
    # mirror of the real table: "Product Name" is a computed column (type 20), "Product" is the writable text column
    f, lc = fake_client(cfg, [("Barcode", 1), ("Product Name", 20), ("Product", 1), ("Quantity", 2), ("Expiration Date", 1)])
    u = add_user(session, "np_admin", "admin")
    np_ = new_products.create(session, u, barcode="990001", name="Test Candy 100g", brand="Acme", flavor="Chocolate", category="Snacks", supplier_name="Acme",
                              unit="case", units_per_pack=24, pack_price=240.0, price=13.0, packs=1, expiry="2027-01-31", note="box of 24", client=lc)
    session.commit()
    assert np_.status == "lark_sent" and np_.lark_record_id == "rec1" and np_.expiry_date.isoformat() == "2027-01-31"
    row = f.rows[0]
    assert row["Product"] == "Test Candy 100g" and "Product Name" not in row  # the computed column is never written
    assert row["Barcode"] == "990001" and row["Selling Price"] == 13.0 and row["Status"] == "Pending RPA"
    assert row["Expiration Date"] == "2027-01-31" and row["Supplier Cost"] == 10.0 and row["Quantity"] == 24  # 240 / 24 per piece; 1 case x 24
    assert row["Brand"] == "Acme" and row["Flavor"] == "Chocolate" and row["Unit"] == "case" and row["Units Per Pack"] == 24 and row["Pack Price"] == 240.0
    assert row["Entry Type"] == "New"
    added = {n for n, t in f.fields}  # missing columns were created once, then written
    assert {"Category", "Supplier", "Supplier Cost", "Selling Price", "Status", "Created By", "Brand", "Flavor", "Unit", "Units Per Pack", "Pack Price", "Entry Type"} <= added
    xlsx = cfg.data_dir / np_.excel_path
    assert xlsx.is_file() and xlsx.parent.name == "new-products"
    from openpyxl import load_workbook
    ws = load_workbook(xlsx).active
    assert [c.value for c in ws[1]] == new_products.EXCEL_HEADERS
    assert ws[2][0].value == "Test Candy 100g" and ws[2][2].value == "990001" and ws[2][8].value == "Pending RPA"
    assert ws[2][5].value == 10.0 and ws[2][7].value == 24 and ws[2][11].value == "New" and ws[2][12].value == "Acme" and ws[2][17].value == 1
    ev = session.scalar(select(AuditEvent).where(AuditEvent.action == "new_product.create"))
    assert ev is not None and (ev.after or {}).get("barcode") == "990001"


def test_lark_failure_keeps_the_registration(cfg, session):
    f, lc = fake_client(cfg, [("Barcode", 1)], fail=True)
    u = add_user(session, "np_admin2", "admin")
    np_ = new_products.create(session, u, barcode="990002", name="Fail Cola", brand="Acme", flavor="Chocolate", cost=5, price=7, qty=1, client=lc)
    session.commit()
    assert np_.status == "lark_failed" and np_.lark_error
    assert (cfg.data_dir / np_.excel_path).is_file()  # the RPA file still exists
    # retry succeeds once Lark is back
    f2, lc2 = fake_client(cfg, [("Barcode", 1), ("Status", 1)])
    new_products.retry_lark(session, u, np_, client=lc2)
    assert np_.status == "lark_sent" and f2.rows[0]["Barcode"] == "990002"


def test_validation_and_existing_barcode(cfg, session, sync, fake):
    seeded(sync, fake, session)  # StoreHub product P1 has barcode 111
    u = add_user(session, "np_admin3", "admin")
    try:
        new_products.create(session, u, barcode="", name="X", brand="Acme", flavor="Chocolate", )
        assert False
    except new_products.NewProductError as e:
        assert "barcode" in str(e)
    try:
        new_products.create(session, u, barcode="111", name="Cola again", brand="Acme", flavor="Chocolate", )
        assert False
    except new_products.NewProductError as e:
        assert "already exists" in str(e)
    try:
        new_products.create(session, u, barcode="990003", name="Neg", brand="Acme", flavor="Chocolate", cost=1, qty=1, price=-1)
        assert False
    except new_products.NewProductError as e:
        assert "negative" in str(e)
    assert session.scalar(select(func.count()).select_from(NewProduct)) == 0



def test_audit_chain_still_verifies(cfg, session):
    from app.audit import verify_chain
    u = add_user(session, "np_admin4", "admin")
    f, lc = fake_client(cfg, [("Barcode", 1)])
    new_products.create(session, u, barcode="990004", name="Chain Check", brand="Acme", flavor="Chocolate", cost=1, price=2, qty=1, client=lc)
    session.commit()
    ok, total, bad = verify_chain(session)
    assert ok and total > 0 and bad is None


# ---------------------------------------------------------------- the product entry flow (scan -> new / existing -> checks -> save)
def _new(session, u, lc, **kw):
    base = dict(barcode="770001", name="Flow Chips 50g", brand="Flowco", flavor="Cheese", unit="case", units_per_pack=24, pack_price=1200.0, price=65.0, packs=2)
    base.update(kw)
    return new_products.create(session, u, client=lc, **base)


def test_cost_price_and_quantity_maths(cfg, session):
    assert new_products.unit_cost(1200, 24) == 50.0 and new_products.selling_price(50, 30) == 65.0
    f, lc = fake_client(cfg, [("Barcode", 1)])
    u = add_user(session, "np_math", "admin")
    np_ = _new(session, u, lc)
    assert np_.cost == 50.0 and np_.qty == 48 and np_.packs == 2 and np_.units_per_pack == 24  # 2 cases x 24 = 48
    np2 = _new(session, u, lc, barcode="770002", name="Loose Candy", flavor="Mint", unit="piece", units_per_pack=None, pack_price=3.0, price=5.0, packs=10)
    assert np2.cost == 3.0 and np2.qty == 10


def test_checks_block_bad_entries(cfg, session):
    f, lc = fake_client(cfg, [("Barcode", 1)])
    u = add_user(session, "np_checks", "admin")
    for kw, word in (({"brand": ""}, "brand"), ({"flavor": ""}, "flavor"), ({"units_per_pack": None}, "how many"), ({"pack_price": None}, "supplier price"),
                     ({"price": 40.0}, "below the cost"), ({"price": None}, "selling price"), ({"packs": None}, "quantity"),
                     ({"packs": 500}, "too many"), ({"packs": 1.5, "units_per_pack": 1.5, "pack_price": 3.0, "price": 4.0}, "whole number"),
                     ({"unit": "pallet"}, "unit")):
        try:
            _new(session, u, lc, **kw)
            assert False, kw
        except new_products.NewProductError as e:
            assert word.lower() in str(e).lower(), (kw, str(e))
    assert session.scalar(select(func.count()).select_from(NewProduct)) == 0
    _new(session, u, lc)
    try:  # same product and flavor under a different barcode
        _new(session, u, lc, barcode="770009")
        assert False
    except new_products.NewProductError as e:
        assert "one barcode per flavor" in str(e)
    _new(session, u, lc, barcode="770010", flavor="Barbecue")  # another flavor is its own record


def test_existing_product_flow(cfg, session, sync, fake):
    seeded(sync, fake, session)  # StoreHub product P1, barcode 111
    f, lc = fake_client(cfg, [("Barcode", 1)])
    u = add_user(session, "np_exist", "admin")
    p = new_products.lookup(session, "111")
    assert p is not None and new_products.lookup(session, "nope") is None
    p.cost, p.unit_price = 10.0, 15.0
    for kw, word in (({"confirmed": False, "packs": 1}, "match the goods"), ({"confirmed": True}, "quantity"), ({"confirmed": True, "packs": 0}, "above zero"),
                     ({"confirmed": True, "packs": 1, "unit": "case", "units_per_pack": 12, "pack_price": 156.0}, "enter the new selling price"),
                     ({"confirmed": True, "packs": 1, "unit": "case", "units_per_pack": 12, "pack_price": 156.0, "price": 12.0}, "below the cost")):
        try:
            new_products.create_existing(session, u, p, client=lc, **kw)
            assert False, kw
        except new_products.NewProductError as e:
            assert word in str(e), (kw, str(e))
    np_ = new_products.create_existing(session, u, p, confirmed=True, unit="case", units_per_pack=12, packs=3, client=lc)  # cost unchanged
    assert np_.entry_type == "existing" and np_.product_id == p.id and np_.qty == 36 and np_.cost == 10.0 and np_.price == 15.0
    np2 = new_products.create_existing(session, u, p, confirmed=True, unit="case", units_per_pack=12, pack_price=156.0, price=18.0, packs=1, client=lc)
    assert np2.cost == 13.0 and np2.price == 18.0 and np2.qty == 12  # the cost changed: cost and price recomputed
    assert f.rows[0]["Entry Type"] == "Existing" and f.rows[1]["Pack Price"] == 156.0
    new_products.create_existing(session, u, p, confirmed=True, unit="piece", packs=5, client=lc)  # an existing product can be entered again and again



def test_internet_lookup_prefills_the_form(cfg, session, sync, fake, client, monkeypatch):
    monkeypatch.setattr(new_products, "web_lookup", _REAL_WEB_LOOKUP)  # this test feeds it a fake internet
    class R:
        status_code = 200
        def json(self): return {"product": {"product_name": "Piattos Cheese 45g", "brands": "Jack-n-Jill, Other", "categories": "Mustard crisps, Snacks", "quantity": "45 g"}}
    calls = []
    monkeypatch.setattr(new_products.httpx, "get", lambda url, **kw: calls.append(url) or R())
    new_products._WEB_CACHE.clear()
    seeded(sync, fake, session)
    g = new_products.web_lookup(session, "4800016644801")
    assert g == {"name": "Piattos Cheese 45g", "brand": "Jack-n-Jill", "category": "Mustard crisps", "quantity": "45 g", "source": "Open Food Facts"}
    new_products.web_lookup(session, "4800016644801")
    assert len(calls) == 1 and "4800016644801" in calls[0]  # cached; only the barcode is in the request
    assert new_products.web_lookup(session, "abc") is None and len(calls) == 1  # junk is never sent
    add_user(session, "np_net", "staff", locs=[2]); login(client, "np_net")
    page = client.get("/new-product?barcode=4800016644801").text  # the page itself never waits for the internet
    assert "Looking it up on the internet" in page and len(calls) == 1
    assert client.get("/new-product/guess?barcode=4800016644801").json()["name"] == "Piattos Cheese 45g"
    new_products._WEB_CACHE.clear()
    monkeypatch.setattr(new_products.httpx, "get", lambda url, **kw: (_ for _ in ()).throw(RuntimeError("offline")))
    assert new_products.web_lookup(session, "4800016644801") is None  # offline: the form just stays empty
    assert client.get("/new-product/guess?barcode=4800016644801").json() == {}



def test_barcode_variants_match_the_same_product(cfg, session, sync, fake):
    seeded(sync, fake, session)
    p = new_products.lookup(session, "111")
    p.barcode = "012345678905"  # the product list holds the 12-digit UPC-A
    session.commit()
    for scanned in ("012345678905", "0012345678905", "12345678905"):  # EAN-13 reading, with and without the leading zeros
        assert new_products.lookup(session, scanned).id == p.id, scanned
    assert new_products.lookup(session, "9999999999999") is None
    u = add_user(session, "np_var", "admin")
    try:
        new_products.create(session, u, barcode="0012345678905", name="X", brand="B", flavor="F", cost=1, price=2, qty=1)
        assert False
    except new_products.NewProductError as e:
        assert "already exists" in str(e)


def test_unknown_barcode_refreshes_the_storehub_list_then_checks_again(cfg, session, sync, fake, monkeypatch):
    seeded(sync, fake, session)
    monkeypatch.setattr(new_products, "products_fresh", lambda db: False)
    calls = []

    def fake_refresh():  # StoreHub answers with a product created there today
        calls.append(1)
        from app.models import Product
        session.add(Product(storehub_id="NEWSH", name="Created in StoreHub today", barcode="5550001"))
        session.commit()
        return True
    monkeypatch.setattr(new_products, "refresh_products", fake_refresh)
    assert new_products.find_product(session, "111") == (new_products.lookup(session, "111"), False) and not calls  # saved list hit: StoreHub not asked
    p, refreshed = new_products.find_product(session, "5550001")
    assert refreshed and p.name == "Created in StoreHub today" and len(calls) == 1
    assert new_products.find_product(session, "000", refresh=False) == (None, False) and len(calls) == 1
    monkeypatch.setattr(new_products, "refresh_products", lambda: False)  # StoreHub unreachable
    assert new_products.find_product(session, "7770002") == (None, False)


def test_refresh_is_throttled(monkeypatch):
    import app.storehub.sync as shs
    runs = []

    class S:
        def sync_products(self):
            runs.append(1)
            return {"received": 1}
    monkeypatch.setattr(shs, "SyncService", S)
    monkeypatch.undo()  # drop the autouse stub so the real function is used
    monkeypatch.setattr(shs, "SyncService", S)
    new_products._last_refresh = 0.0
    assert new_products.refresh_products() is True and new_products.refresh_products() is False and len(runs) == 1


def test_scan_page_is_ready_for_a_handheld_scanner(cfg, session, sync, fake, client):
    seeded(sync, fake, session)
    add_user(session, "np_hw", "staff", locs=[2]); login(client, "np_hw")
    for url in ("/new-product", "/new-product?barcode=424242", "/new-product?barcode=111"):  # ready on every step, so the next product can be scanned at once
        page = client.get(url).text
        assert "Check barcode" not in page and 'aria-label="Scan"' in page  # no button: it checks by itself; the scan button is an icon
        assert "data-scanner" in page and "/static/scanner.js" in page and 'data-scan="np-scan"' in page and "/static/scan.js" in page, url
    js = client.get("/static/scanner.js")
    assert js.status_code == 200 and "keydown" in js.text
    cam = client.get("/static/scan.js")
    assert cam.status_code == 200 and "getUserMedia" in cam.text and client.get("/static/vendor/zxing-library.min.js").status_code == 200


def test_a_fresh_product_list_is_not_fetched_again(cfg, session, sync, fake, monkeypatch):
    seeded(sync, fake, session)  # the seeding sync just fetched the product list from StoreHub
    calls = []
    monkeypatch.setattr(new_products, "refresh_products", lambda: calls.append(1) or True)
    assert new_products.products_fresh(session) is True
    assert new_products.find_product(session, "7770003") == (None, False) and not calls  # the answer would be the same: no wait
