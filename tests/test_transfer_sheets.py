"""Warehouse transfer sheet: scan + quantity instead of the notebook; saved with an Excel/CSV for StoreHub; matched back to StoreHub's transfer."""
from datetime import timedelta

import pytest
from sqlalchemy import select

from app import audit
from app.models import AuditEvent, Location, Product, StockTransfer, StockTransferImport, StockTransferLine, TransferSheet
from app.services import transfer_sheets as ts
from app.timeutil import now_utc
from tests.test_services import mkuser, seeded
from tests.test_web import add_user, client, csrf_of, login  # noqa: F401


def _store(session):
    return session.scalar(select(Location).where(Location.storehub_id == "S_A"))


def _p(session, sid):
    return session.scalar(select(Product).where(Product.storehub_id == sid))


def test_scan_add_merge_edit_remove(sync, fake, session):
    seeded(sync, fake, session)
    wh = mkuser(session, "wh1", "warehouse")
    sheet = ts.create(session, wh, _store(session).id); session.commit()
    assert sheet.source.type == "warehouse" and sheet.target.name.endswith("Alpha") and sheet.status == "draft"
    assert ts.create(session, wh, _store(session).id).id == sheet.id  # an open draft for the same store is reused
    exact, many = ts.find(session, "111")  # a barcode scan is one exact product
    assert exact.name == "Cola 500ml" and many == []
    assert ts.find(session, "CH1")[0].name == "Chips"  # SKU works too
    exact, many = ts.find(session, "chip")  # a typed name is searched
    assert exact is None and [p.name for p in many] == ["Chips"]
    assert ts.find(session, "zzz-nothing") == (None, [])
    line = ts.add_line(session, sheet, exact_p := _p(session, "P1"), 24); session.commit()
    ts.add_line(session, sheet, exact_p, 12)  # scanning the same product again adds to its line
    assert len(sheet.lines) == 1 and sheet.lines[0].qty == 36 and sheet.lines[0].unit_cost == 12.0
    for bad in (0, -3, "x", 2.5, 10 ** 7):
        with pytest.raises(ts.SheetError):
            ts.add_line(session, sheet, exact_p, bad)
    ts.add_line(session, sheet, _p(session, "P2"), 6)
    ts.set_qty(session, sheet, line.id, 30)
    assert {l.product.name: l.qty for l in sheet.lines} == {"Cola 500ml": 30, "Chips": 6}
    assert ts.on_hand(session, sheet.source_location_id, [exact_p.id]) == {exact_p.id: 100.0}  # the warehouse stock shown next to the line
    ts.remove_line(session, sheet, line.id)
    assert [l.product.name for l in sheet.lines] == ["Chips"]


def test_save_writes_excel_csv_and_audit_then_locks(cfg, sync, fake, session):
    seeded(sync, fake, session)
    wh = mkuser(session, "wh2", "warehouse")
    sheet = ts.create(session, wh, _store(session).id)
    with pytest.raises(ts.SheetError, match="at least one"):
        ts.submit(session, wh, sheet)
    ts.add_line(session, sheet, _p(session, "P1"), 24); ts.add_line(session, sheet, _p(session, "P2"), 12)
    ts.submit(session, wh, sheet); session.commit()
    assert sheet.status == "submitted" and sheet.submitted_at
    from openpyxl import load_workbook
    path = cfg.data_dir / sheet.excel_path
    assert path.parent.name == "stock-transfers" and path.is_file()
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == ts.EXCEL_HEADER
    row = [c.value for c in ws[2]]
    assert row[0] == 'Cola 500ml' and row[2] == 'P1' and row[3] == 24 and row[6] == '111' and row[7].startswith('TheDailyCentro Warehouse') and row[8].endswith('Alpha')
    csv_text = ts.csv_text(sheet)
    assert csv_text.startswith('\ufeffProduct Name,SKU,Product Id,Ordered Qty,Unit,Serial Numbers\r\n#Optional,') and 'Cola 500ml,,P1,24,,' in csv_text
    assert 'Chips,CH1,P2,12,,' in csv_text and ts.file_stem(sheet).startswith('StockTransfer-TS-')
    ev = session.scalar(select(AuditEvent).where(AuditEvent.action == "transfer_sheet.submit"))
    assert ev.after["qty"] == 36 and ev.after["lines"] == 2 and audit.verify_chain(session)[0] is True
    with pytest.raises(ts.SheetError, match="already submitted"):
        ts.add_line(session, sheet, _p(session, "P1"), 1)  # a saved sheet is a record: no more edits


def test_storehub_number_and_automatic_match(sync, fake, session):
    seeded(sync, fake, session)
    wh = mkuser(session, "wh3", "warehouse")
    sheet = ts.create(session, wh, _store(session).id)
    ts.add_line(session, sheet, _p(session, "P1"), 24); ts.add_line(session, sheet, _p(session, "P2"), 12)
    with pytest.raises(ts.SheetError, match="Save the sheet first"):
        ts.mark_in_storehub(session, wh, sheet, "ST1")
    ts.submit(session, wh, sheet); session.commit()
    imp = StockTransferImport(token="t" * 8, filename="x.csv", sha256="0" * 64, raw_csv=b"", created_by=wh.id)
    session.add(imp); session.flush()

    def transfer(ref, qty2, status="Created"):
        t = StockTransfer(reference=ref, source_name="W", target_name="A", source_location_id=sheet.source_location_id, target_location_id=sheet.target_location_id,
                          status=status, created_sh=now_utc(), payload_hash=ref, last_import_id=imp.id)
        t.lines = [StockTransferLine(line_no=1, product_id=_p(session, "P1").id, product_name="Cola", quantity=24),
                   StockTransferLine(line_no=2, product_id=_p(session, "P2").id, product_name="Chips", quantity=qty2)]
        session.add(t)
    transfer("ST0100", 10)  # different quantity: not the same goods
    transfer("ST0101", 12, "Cancelled")  # cancelled in StoreHub: ignored
    session.commit()
    assert ts.match_storehub(session) == 0 and sheet.status == "submitted"
    transfer("ST0102", 12); session.commit()
    assert ts.match_storehub(session) == 1 and sheet.status == "in_storehub" and sheet.sh_reference == "ST0102"
    with pytest.raises(ts.SheetError, match="StoreHub transfer number"):
        ts.mark_in_storehub(session, wh, sheet, "  ")
    ts.mark_in_storehub(session, wh, sheet, " st0103 ")  # typed by hand
    assert sheet.sh_reference == "ST0103"


def test_pages_scan_flow_and_permissions(cfg, sync, fake, session, client):
    seeded(sync, fake, session)
    add_user(session, "wh_web", "warehouse"); add_user(session, "store_web", "staff", locs=[_store(session).id])
    login(client, "store_web")
    assert client.get("/transfer-sheets").status_code == 403  # stores confirm receipt; they do not build warehouse sheets
    client.post("/logout")
    client2 = client
    login(client2, "wh_web")
    page = client2.get("/transfer-sheets")
    assert page.status_code == 200 and "Alpha" in page.text
    csrf = csrf_of(client2, "/transfer-sheets")
    r = client2.post("/transfer-sheets/new", data={"csrf": csrf, "target_id": str(_store(session).id)})
    assert r.status_code == 303
    sid = int(r.headers["location"].rsplit("/", 1)[1])
    sheet_page = client2.get(f"/transfer-sheets/{sid}").text
    assert "data-scanner" in sheet_page and 'data-scan="ts-scan"' in sheet_page and "/static/scanner.js" in sheet_page
    f = client2.get(f"/transfer-sheets/{sid}/find", params={"code": "111"}).json()
    assert f["exact"]["name"] == "Cola 500ml" and f["exact"]["on_hand"] == 100.0 and f["exact"]["on_sheet"] == 0
    assert client2.get(f"/transfer-sheets/{sid}/find", params={"code": "chip"}).json()["matches"][0]["name"] == "Chips"
    assert client2.get(f"/transfer-sheets/{sid}/find", params={"code": "nope"}).json() == {"exact": None, "matches": [], "code": "nope"}
    pid = f["exact"]["id"]
    bad = client2.post(f"/transfer-sheets/{sid}/add", data={"csrf": csrf, "product_id": pid, "qty": "0"})
    assert bad.status_code == 400 and bad.json()["ok"] is False
    assert client2.post(f"/transfer-sheets/{sid}/add", data={"csrf": "wrong", "product_id": pid, "qty": "5"}).status_code in (400, 403)
    ok = client2.post(f"/transfer-sheets/{sid}/add", data={"csrf": csrf, "product_id": pid, "qty": "24"}).json()
    assert ok["ok"] and ok["total_lines"] == 1 and ok["total_qty"] == 24 and ok["lines"][0]["name"] == "Cola 500ml"
    lid = ok["lines"][0]["id"]
    assert client2.post(f"/transfer-sheets/{sid}/line/{lid}", data={"csrf": csrf, "qty": "30"}).json()["total_qty"] == 30
    assert client2.get(f"/transfer-sheets/{sid}/download.csv").status_code == 404  # a draft has no file yet
    assert client2.post(f"/transfer-sheets/{sid}/submit", data={"csrf": csrf}).status_code == 303
    assert client2.get(f"/transfer-sheets/{sid}/download.csv").status_code == 200 and b"Cola 500ml" in client2.get(f"/transfer-sheets/{sid}/download.csv").content
    x = client2.get(f"/transfer-sheets/{sid}/download.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert client2.post(f"/transfer-sheets/{sid}/storehub", data={"csrf": csrf, "reference": "ST0200"}).status_code == 303
    session.expire_all()
    assert session.get(TransferSheet, sid).sh_reference == "ST0200"
    assert "ST0200" in client2.get("/transfer-sheets").text


from pathlib import Path  # noqa: E402

R_DOCS = Path(__file__).resolve().parent.parent / "docs" / "storehub"


def _template_rows():
    import csv
    import io
    raw = (R_DOCS / "StoreHub Stock Transfer Items Import Template.csv").read_text(encoding="utf-8-sig")
    return list(csv.reader(io.StringIO(raw)))


def test_import_file_follows_storehubs_own_template(sync, fake, session):
    """The header and the note row are exactly the ones in the template downloaded from the StoreHub back office."""
    import csv
    import io
    seeded(sync, fake, session)
    tpl = _template_rows()
    assert ts.IMPORT_HEADER == tpl[0] and ts.IMPORT_NOTES == tpl[1]
    wh = mkuser(session, "wh_tpl", "warehouse")
    sheet = ts.create(session, wh, _store(session).id)
    ts.add_line(session, sheet, _p(session, "P1"), 24); ts.add_line(session, sheet, _p(session, "P2"), 6)
    ts.submit(session, wh, sheet); session.commit()
    raw = ts.csv_text(sheet)
    assert raw.startswith("﻿") and "\r\n" in raw and "\n" not in raw.replace("\r\n", "")  # UTF-8 with BOM, CRLF line ends, like the template
    rows = list(csv.reader(io.StringIO(raw.lstrip("﻿"))))
    assert rows[0] == tpl[0] and rows[1] == tpl[1]
    assert rows[2:] == [["Cola 500ml", "", "P1", "24", "", ""], ["Chips", "CH1", "P2", "6", "", ""]]  # Product Name, SKU, Product Id, Ordered Qty, Unit, Serial
    assert all(len(r) == 6 for r in rows)


def test_a_product_with_several_barcodes_is_found_by_any_of_them(sync, fake, session):
    from app.services import new_products
    seeded(sync, fake, session)
    p = _p(session, "P2")
    p.barcode = "100070161773,6977318633319"
    session.commit()
    for code in ("100070161773", "6977318633319", "0006977318633319"):
        assert new_products.lookup(session, code).id == p.id, code
    assert ts.find(session, "6977318633319")[0].id == p.id
    assert new_products.lookup(session, "697731863331") is None  # a part of a barcode is not a match


def test_breadcrumbs_lead_back_up_so_the_browser_back_button_is_not_needed(cfg, sync, fake, session, client):
    import re
    seeded(sync, fake, session)
    add_user(session, "wh_crumb", "warehouse"); login(client, "wh_crumb")

    def trail(url):
        html = client.get(url).text
        m = re.search(r'<header class="topbar">.*?<nav class="crumbs[^>]*>(.*?)</nav>', html, re.S)
        return m and (re.findall(r'<a href="([^"]+)">', m.group(1)), re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", m.group(1))).strip().replace("›", ">"))
    assert trail("/") is None  # the home page shows its own title in the bar, no trail
    links, text = trail("/transfer-sheets")
    assert links == ["/"] and text.endswith("> Transfer sheet")  # Today > Transfer sheet
    csrf = csrf_of(client, "/transfer-sheets")
    sid = client.post("/transfer-sheets/new", data={"csrf": csrf, "target_id": str(_store(session).id)}).headers["location"].rsplit("/", 1)[1]
    links, text = trail(f"/transfer-sheets/{sid}")
    assert links == ["/", "/transfer-sheets"] and text.endswith("Transfer sheet > TS-%05d" % int(sid))  # Today > Transfer sheet > TS-00001
    links, text = trail("/replenishment/repairs")
    assert links == ["/"] and text.endswith("Repair queue")  # the longest sidebar entry wins over /replenishment
    links, text = trail("/replenishment/batches")
    assert links == ["/", "/replenishment"] and text.endswith("> Batches")




def test_a_transfer_only_takes_products_storehub_has(cfg, sync, fake, session, client):
    seeded(sync, fake, session)
    add_user(session, "wh_cat", "warehouse"); login(client, "wh_cat")
    csrf = csrf_of(client, "/transfer-sheets")
    sid = int(client.post("/transfer-sheets/new", data={"csrf": csrf, "target_id": str(_store(session).id)}).headers["location"].rsplit("/", 1)[1])
    cat = client.get("/transfer-sheets/catalog.json").json()
    row = next(r for r in cat if r[1] == "Cola 500ml")
    assert row == [row[0], "Cola 500ml", "", ["111"], 100.0]  # id, name, sku, barcodes, warehouse stock: all the page needs to answer a scan itself
    for fields in ({"product_id": "0", "qty": "4"}, {"product_id": "999999", "qty": "4"}):  # a barcode StoreHub does not know has no product to add
        bad = client.post(f"/transfer-sheets/{sid}/add", data={"csrf": csrf, **fields})
        assert bad.status_code == 400 and bad.json()["ok"] is False
    ok = client.post(f"/transfer-sheets/{sid}/add", data={"csrf": csrf, "product_id": str(row[0]), "qty": "4"}).json()
    assert ok["ok"] and ok["total_qty"] == 4 and ok["lines"][0]["name"] == "Cola 500ml"
    page = client.get(f"/transfer-sheets/{sid}").text
    assert "Add anyway" not in page and 'id="tm-name"' not in page and 'id="ts-quick"' in page and "/new-product?barcode=" not in page  # unknown barcode: the create-in-StoreHub panel opens right on the sheet
    assert client.get(f"/transfer-sheets/{sid}/not-in-storehub.csv").status_code == 404


def test_a_transfer_sheet_can_be_deleted_for_good(cfg, sync, fake, session, client):
    from sqlalchemy import func
    from app.models import TransferSheetLine
    seeded(sync, fake, session)
    add_user(session, "wh_del", "warehouse"); login(client, "wh_del")
    csrf = csrf_of(client, "/transfer-sheets")
    sid = int(client.post("/transfer-sheets/new", data={"csrf": csrf, "target_id": str(_store(session).id)}).headers["location"].rsplit("/", 1)[1])
    pid = _p(session, "P1").id
    client.post(f"/transfer-sheets/{sid}/add", data={"csrf": csrf, "product_id": str(pid), "qty": "12"})
    client.post(f"/transfer-sheets/{sid}/submit", data={"csrf": csrf})
    session.expire_all()
    xl = cfg.data_dir / session.get(TransferSheet, sid).excel_path
    assert xl.is_file() and "/delete" in client.get("/transfer-sheets").text and "/delete" in client.get(f"/transfer-sheets/{sid}").text
    assert client.post(f"/transfer-sheets/{sid}/delete", data={"csrf": "wrong"}).status_code in (400, 403)
    assert session.get(TransferSheet, sid) is not None
    r = client.post(f"/transfer-sheets/{sid}/delete", data={"csrf": csrf})
    assert r.status_code == 303 and "err=" not in r.headers["location"]
    session.expire_all()
    assert session.get(TransferSheet, sid) is None and not xl.exists()
    assert session.scalar(select(func.count()).select_from(TransferSheetLine).where(TransferSheetLine.sheet_id == sid)) == 0
    ev = session.scalar(select(AuditEvent).where(AuditEvent.action == "transfer_sheet.delete"))
    assert ev.before["lines"] == [{"barcode": "111", "name": "Cola 500ml", "qty": 12.0}] and audit.verify_chain(session)[0] is True  # the history of what it held stays in the audit log
    assert client.post(f"/transfer-sheets/{sid}/delete", data={"csrf": csrf}).status_code == 404
    # a draft can be deleted too
    sid2 = int(client.post("/transfer-sheets/new", data={"csrf": csrf, "target_id": str(_store(session).id)}).headers["location"].rsplit("/", 1)[1])
    assert client.post(f"/transfer-sheets/{sid2}/delete", data={"csrf": csrf}).status_code == 303


def test_discarding_a_sheet_deletes_it_so_numbers_restart_from_what_exists(sync, fake, session):
    seeded(sync, fake, session)
    wh = mkuser(session, "wh_num", "warehouse")
    first = ts.create(session, wh, _store(session).id); session.commit()
    ts.add_line(session, first, _p(session, "P1"), 5); session.commit()
    assert first.id == 1
    ts.cancel(session, wh, first); session.commit()
    assert session.scalars(select(TransferSheet)).all() == []  # gone, not hidden
    assert session.scalar(select(AuditEvent).where(AuditEvent.action == "transfer_sheet.cancel")) is not None  # the audit log still says it existed
    again = ts.create(session, wh, _store(session).id); session.commit()
    assert again.id == 1  # nothing left, so it is TS-00001 again
    ts.add_line(session, again, _p(session, "P1"), 5); ts.submit(session, wh, again); session.commit()
    nxt = ts.create(session, wh, _store(session).id); session.commit()
    assert nxt.id == 2  # carries on after a saved sheet


def test_quick_product_from_the_sheet_goes_straight_on(cfg, sync, fake, session, client, monkeypatch):
    """Scan an unknown barcode on the sheet: the panel creates it in StoreHub and it lands on the sheet in the same flow."""
    from app.services import product_push
    from tests.test_product_push import FakeCreateStoreHub
    seeded(sync, fake, session)
    add_user(session, "ts_qp", "warehouse")
    sh = FakeCreateStoreHub()
    monkeypatch.setattr(product_push, "StoreHubClient", lambda: sh.client())
    login(client, "ts_qp")
    csrf = csrf_of(client, "/transfer-sheets")
    r = client.post("/transfer-sheets/new", data={"csrf": csrf, "target_id": str(_store(session).id)})
    assert r.status_code == 303
    sid = int(r.headers["location"].rsplit("/", 1)[1])
    page = client.get(f"/transfer-sheets/{sid}").text
    assert 'id="ts-quick"' in page and "Add to StoreHub" in page  # the panel is on the page, wired by the page script
    r = client.post(f"/transfer-sheets/{sid}/quick-product",
                    data={"csrf": csrf, "barcode": "770009", "name": "Sheet Quick Candy", "category": "Snacks", "price": "25", "cost": "15"})
    assert r.status_code == 200 and r.json()["ok"], r.text
    pid, name, sku, barcodes, on_hand = r.json()["product"]
    assert name == "Sheet Quick Candy" and barcodes == ["770009"] and on_hand is None
    a = client.post(f"/transfer-sheets/{sid}/add", data={"csrf": csrf, "product_id": pid, "qty": "6"})
    assert a.json()["ok"] and a.json()["lines"][0]["name"] == "Sheet Quick Candy" and a.json()["lines"][0]["qty"] == 6
    dup = client.post(f"/transfer-sheets/{sid}/quick-product", data={"csrf": csrf, "barcode": "770009", "name": "Twin", "price": "5"})
    assert dup.status_code == 400 and "Already in StoreHub" in dup.json()["error"]
    assert len(sh.posts) == 1  # the duplicate never reached StoreHub
    # a store user cannot use it (stores confirm receipt; they do not build sheets or create products)
    add_user(session, "ts_qp_store", "staff", locs=[_store(session).id])
    client.post("/logout")
    login(client, "ts_qp_store")
    assert client.post(f"/transfer-sheets/{sid}/quick-product",
                       data={"csrf": csrf, "barcode": "770010", "name": "Nope", "price": "5"}).status_code == 403
