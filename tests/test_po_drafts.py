"""Write a purchase order: supplier + delivery first, products with exact quantity x price (no 2-decimal limit), save -> StoreHub CSV + Excel + print."""
import csv
import io
import re
from decimal import Decimal

import pytest
from sqlalchemy import select

from app.models import Location, PoDraft, Product, Supplier
from app.services import po_drafts as pod
from tests.test_web import add_user, client, csrf_of, login, seeded  # noqa: F401 (fixtures)

BASE = "/purchase-orders/new"


def _setup(sync, fake, session):
    seeded(sync, fake, session)
    session.add(Supplier(name="Basic Essential Marketing Inc.", name_source="manual"))
    session.commit()
    add_user(session, "po_w", "warehouse")
    loc = session.scalar(select(Location).where(Location.type == "warehouse"))
    return loc, session.scalar(select(Product).where(Product.in_api.is_(True)))


def _start(client, session, loc):
    csrf = csrf_of(client, BASE)
    r = client.post(BASE + "/start", data={"csrf": csrf, "supplier": "basic essential marketing inc."})
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    return csrf


def test_totals_are_exact_and_never_cut_to_two_decimals():
    assert pod.total_of("12", "254.4642857") == "3053.5714284"
    assert pod.total_of("0.5", "0.333333333") == "0.1666666665"
    assert pod.total_of("3", "10") == "30"
    assert pod.show("3053.5714284") == "3,053.5714284" and pod.show("18408") == "18,408.00" and pod.show("0.5") == "0.50"
    for bad in ("", "abc", "-1", "1.2.3", "1e5"):
        with pytest.raises(pod.PoError):
            pod.check_price(bad)
    assert pod.check_price("1,234.50") == "1234.5" and pod.check_price(".25") == "0.25"
    # StoreHub keeps a supplier price to 6 decimals and cuts the rest off (it shows 254.464285 for 285 / 1.12): so does every price here, never rounded up
    assert pod.check_price("254.4642857") == "254.464285" and pod.check_price("0.9999999") == "0.999999" and pod.check_price("10.123456") == "10.123456"
    assert pod.check_price("3.1234567891234567") == "3.123456"
    with pytest.raises(pod.PoError):
        pod.check_qty("0")


def test_choose_supplier_and_delivery_first_then_add_products_and_save(cfg, sync, fake, session, client):
    loc, product = _setup(sync, fake, session)
    login(client, "po_w")
    page = client.get(BASE)
    assert page.status_code == 200 and "Basic Essential Marketing Inc." in page.text and 'name="supplier"' in page.text
    assert 'id="po-scan"' not in page.text  # no product box until the supplier and where to deliver are chosen
    csrf = csrf_of(client, BASE)
    r = client.post(BASE + "/start", data={"csrf": csrf, "supplier": "Nobody Ltd"})
    assert "err=" in r.headers["location"]
    csrf = _start(client, session, loc)
    assert 'id="po-scan"' in client.get(BASE).text
    r = client.post(BASE + "/add", data={"csrf": csrf, "product_id": product.id, "qty": "12", "price": "254.4642857"})
    assert r.status_code == 200 and r.json()["total_lines"] == 1
    line = r.json()["lines"][0]
    assert line["qty"] == "12" and line["price"] == "254.464285" and line["total"] == "3053.57142"  # the 7th decimal is cut off, like StoreHub
    # subtotal, 12% tax on it, and the overall total = subtotal + tax: exact, never cut to two decimals
    assert r.json()["subtotal"] == "3053.57142" and r.json()["tax"] == "366.4285704" and r.json()["total"] == "3419.9999904" and r.json()["rate"] == "12"
    # the same product again adds to its line and takes the new price
    r = client.post(BASE + "/add", data={"csrf": csrf, "product_id": product.id, "qty": "0.5", "price": "260.123456789"}).json()
    assert r["total_lines"] == 1 and r["lines"][0]["qty"] == "12.5" and r["lines"][0]["price"] == "260.123456" and r["lines"][0]["total"] == "3251.5432"
    lid = r["lines"][0]["id"]
    r = client.post(f"{BASE}/line/{lid}", data={"csrf": csrf, "qty": "10", "price": "1.123456"}).json()
    assert r["ok"] and r["lines"][0]["total"] == "11.23456"
    assert r["subtotal"] == "11.23456" and r["tax"] == "1.3481472" and r["total"] == "12.5827072"
    bad = client.post(f"{BASE}/line/{lid}", data={"csrf": csrf, "qty": "x", "price": "1"})
    assert bad.status_code == 400 and bad.json()["lines"][0]["qty"] == "10"  # the line is untouched and the screen is sent the true values
    save = client.post(BASE + "/save", data={"csrf": csrf})
    assert save.status_code == 303 and "err=" not in save.headers["location"]
    session.expire_all()
    d = session.scalar(select(PoDraft))
    assert d.status == "saved" and d.location_id == loc.id and d.excel_path and d.supplier.name == "Basic Essential Marketing Inc."
    csv_resp = client.get(f"{BASE}/lists/{d.id}/download.csv")
    rows = list(csv.reader(io.StringIO(csv_resp.content.decode("utf-8-sig"))))
    assert rows[0] == ["Product Name", "SKU", "Product Id", "Qty to Order ", "Unit", "Supplier Price"]
    assert rows[1] == [product.name, product.sku or "", product.storehub_id, "10", "", "1.123456"]
    xl = client.get(f"{BASE}/lists/{d.id}/download.xlsx")
    assert xl.status_code == 200 and xl.content[:2] == b"PK"
    from openpyxl import load_workbook
    sheet = list(load_workbook(io.BytesIO(xl.content)).active.iter_rows(values_only=True))
    assert sheet[0] == tuple(pod.EXCEL_HEADER) and sheet[1][5] == 1.123456 and sheet[1][6] == 11.23456 and sheet[1][8] == "Basic Essential Marketing Inc."
    assert (cfg.data_dir / d.excel_path).parent == (cfg.data_dir / pod.RPA_DIR)
    printed = client.get(f"{BASE}/lists/{d.id}/print")
    assert printed.status_code == 200 and "Purchase Order" in printed.text and "1.123456" in printed.text and product.name in printed.text
    assert "11.23456" in printed.text  # the printout keeps every decimal of the amount and the total (not cents)
    assert "1.35" in printed.text and "12.582707" in printed.text  # the tax row and the overall total (subtotal + tax)
    sheet_rows = [r for r in sheet if r[0] in ("Subtotal", "Tax 12%", "Total")]
    assert [r[6] for r in sheet_rows] == [11.23456, 1.3481472, 12.5827072]
    assert client.get(f"{BASE}/lists/{d.id}").status_code == 200
    # a saved order cannot be changed; the next one starts fresh
    assert client.post(BASE + "/add", data={"csrf": csrf, "product_id": product.id, "qty": "1", "price": "1"}).status_code == 404
    assert client.post(f"{BASE}/lists/{d.id}/delete", data={"csrf": csrf}).status_code == 303
    assert client.get(f"{BASE}/lists/{d.id}/download.csv").status_code == 404


def test_validation_discard_and_privacy(sync, fake, session, client):
    loc, product = _setup(sync, fake, session)
    add_user(session, "po_other", "warehouse")
    login(client, "po_w")
    csrf = _start(client, session, loc)
    for bad, why in (({"qty": "0"}, "above zero"), ({"qty": ""}, "quantity"), ({"price": ""}, "supplier price"), ({"price": "abc"}, "number"), ({"product_id": 999999}, "StoreHub")):
        r = client.post(BASE + "/add", data={"csrf": csrf, "product_id": product.id, "qty": "2", "price": "3", **bad})
        assert r.status_code == 400 and why in r.json()["error"], (bad, r.text)
    assert "err=" in client.post(BASE + "/save", data={"csrf": csrf}).headers["location"]  # nothing on the order yet
    assert client.post(BASE + "/add", data={"csrf": csrf, "product_id": product.id, "qty": "2", "price": "3"}).json()["ok"]
    did = session.scalar(select(PoDraft.id))
    client.post(BASE + "/save", data={"csrf": csrf})
    client.cookies.clear()
    login(client, "po_other")
    assert client.get(f"{BASE}/lists/{did}").status_code == 404 and client.get(f"{BASE}/lists/{did}/download.csv").status_code == 404
    csrf = csrf_of(client, BASE)
    _start(client, session, loc)
    client.post(BASE + "/add", data={"csrf": csrf, "product_id": product.id, "qty": "1", "price": "1"})
    assert client.post(BASE + "/discard", data={"csrf": csrf}).status_code == 303
    session.expire_all()
    assert len(session.scalars(select(PoDraft)).all()) == 1  # only the other person's saved order is left


def test_store_roles_cannot_open_it(sync, fake, session, client):
    _setup(sync, fake, session)
    add_user(session, "po_staff", "staff", locs=[2])
    login(client, "po_staff")
    assert client.get(BASE).status_code == 403


def test_the_page_is_reachable_as_new_and_not_read_as_a_po_number(sync, fake, session, client):
    _setup(sync, fake, session)
    login(client, "po_w")
    assert client.get(BASE).status_code == 200
    assert client.get("/purchase-orders").status_code == 200 and "/purchase-orders/new" in client.get("/purchase-orders").text


def test_breadcrumb_steps_never_link_to_a_missing_page(client):
    from types import SimpleNamespace
    from app.web.deps import _crumbs
    c = SimpleNamespace(t=lambda x: x, request=SimpleNamespace(app=client.app))
    assert _crumbs(c, "/sales/tx/6090", None) == [("Dashboard", "/"), ("Sales", "/sales"), ("Transactions", "/sales/transactions"), ("#6090", None)]
    assert _crumbs(c, "/sales/product/12", "Cola")[-2] == ("Sales by product", "/sales/products")
    trail = _crumbs(c, "/purchase-orders/new/lists/3", "RO-00003")
    assert ("Saved orders", "/purchase-orders/new?saved=1") in trail and trail[-1] == ("RO-00003", None)
    for label, url in _crumbs(c, "/sales/transactions", None) + _crumbs(c, "/purchase-orders/7", None):
        assert url is None or url == "/" or any(r.path_regex.match(url.split("?")[0]) for r in client.app.router.routes if hasattr(r, "path_regex")), (label, url)


def test_a_storehub_price_is_shown_cut_off_after_six_decimals():
    from app.web.deps import price6
    assert price6(254.4642857142857) == "254.464285" and price6(227.67857142857142) == "227.678571" and price6(189.28571428571428) == "189.285714"
    assert price6(12) == "12.00" and price6(10.125) == "10.125" and price6(1234.5) == "1,234.50" and price6(None) == ""
