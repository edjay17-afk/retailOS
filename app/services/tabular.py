"""Reading CSV and Excel (.xlsx) uploads: rows, column headings, products and locations. Shared by the expiry, stock count and stock change imports."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Location, Product

MAX_ROWS = 5000


def norm(v) -> str:
    return re.sub(r"[^a-z0-9 ]", "", re.sub(r"[_\-/]+", " ", str(v or "").strip().lower())).strip()


def read_table(filename: str, data: bytes) -> list[list]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ValueError("Could not read this Excel file. Save it as .xlsx or .csv and try again")
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported. In Excel choose Save As, then .xlsx or CSV")
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))]


def to_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()
    if not t:
        return None
    try:
        return date.fromisoformat(t[:10]) if re.match(r"^\d{4}-\d{1,2}-\d{1,2}", t) else _slash_date(t)
    except ValueError:
        return None


def _slash_date(t: str) -> date:
    m = re.match(r"^(\d{1,2})[/.](\d{1,2})[/.](\d{2,4})$", t)
    if m:
        a, b, y = int(m[1]), int(m[2]), int(m[3])
        y += 2000 if y < 100 else 0
        return date(y, b, a) if a > 12 else date(y, a, b)  # 31/12/2026 is day first; 03/04/2026 is read month first
    for fmt in ("%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%B %d, %Y", "%d-%b-%y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    raise ValueError(t)


def to_qty(v) -> float | None:
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def find_product(db: Session, key: str) -> tuple[Product | None, str]:
    key = key.strip()
    if re.fullmatch(r"\d+\.0", key):  # Excel turns barcodes into numbers
        key = key[:-2]
    exact = list(db.scalars(select(Product).where((Product.barcode == key) | (Product.sku == key)).limit(2)))
    if len(exact) == 1: return exact[0], ""
    if len(exact) > 1: return None, "Barcode matches multiple accounts; select a single POS account first"
    found = list(db.scalars(select(Product).where(Product.in_api.is_(True), Product.name == key).limit(2)))
    if not found:
        found = list(db.scalars(select(Product).where(Product.in_api.is_(True), Product.name.like(f"%{key}%")).limit(3)))
    if len(found) == 1:
        return found[0], ""
    return None, ("no product matches" if not found else "more than one product matches, use the barcode")


def map_columns(head: list, fields: dict[str, tuple[str, ...]], required: tuple[str, ...]) -> dict[str, int]:
    """Column positions by heading: any of the accepted names, in any order. Raises a readable error when a required column is missing."""
    heads = [norm(x) for x in head]
    col: dict[str, int] = {}
    for field, names in fields.items():
        for i, h in enumerate(heads):
            if h in names and i not in col.values():
                col[field] = i
                break
    missing = [f for f in required if f not in col]
    if missing:
        raise ValueError("The first row must name the columns. Missing: " + ", ".join(missing) + ". Download the template to see the layout")
    return col


def body(table: list[list], limit: int = MAX_ROWS) -> list[list]:
    rows = table[1:]
    if len(rows) > limit:
        raise ValueError(f"At most {limit} rows per file")
    return rows


def match_location(value, locations: list[Location]) -> Location | None:
    """A store named in a file: exact or partial, ignoring case and the 'The Daily Centro -' prefix."""
    k = norm(value)
    if not k:
        return None
    by_name = {norm(l.name): l for l in locations}
    return by_name.get(k) or next((l for nm, l in by_name.items() if k in nm or nm.replace("the daily centro", "").strip() == k), None)


def cell(row: list, col: dict[str, int], field: str):
    i = col.get(field)
    return row[i] if i is not None and i < len(row) else None


def clean_rows(table: list[list]) -> list[list]:
    return [r for r in table if r and any(str(x or "").strip() for x in r)]
